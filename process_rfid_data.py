#!/usr/bin/env python3
"""
RFID Data Processing Script

This script processes RFID tag data from a text file and organizes it into an Excel spreadsheet.
The output Excel file will have:
- First column: Tag names (EPC)
- First row: Frequencies
- Second row: RSSI, Loss Rate, Initial Phase, Final Phase data for each frequency

Usage: python process_rfid_data.py <input_file.txt>
"""

import re
import sys
import pandas as pd
from collections import defaultdict


def parse_rfid_data(file_path):
    """
    Parse RFID data from text file and organize by tag and frequency

    Args:
        file_path (str): Path to the input text file

    Returns:
        dict: Dictionary with tag EPC as keys and frequency data as values
    """
    tag_data = defaultdict(dict)

    with open(file_path, 'r', encoding='utf-8') as file:
        lines = file.readlines()

    current_tag = None
    current_data = {}

    for line in lines:
        line = line.strip()

        # Check for tag EPC line
        if line.startswith("Tag EPC:"):
            # Save previous tag data if exists
            if current_tag and current_data:
                freq = current_data.get('frequency')
                if freq:
                    tag_data[current_tag][freq] = current_data.copy()

            # Extract tag EPC
            current_tag = line.split("Tag EPC:")[1].strip()
            current_data = {}

        # Extract data fields
        elif line.startswith("Reads:"):
            reads_match = re.search(r'Reads:\s*(\d+)/(\d+)', line)
            if reads_match:
                current_data['reads'] = int(reads_match.group(1))
                current_data['total_reads'] = int(reads_match.group(2))

        elif line.startswith("Packet Loss Rate:"):
            loss_match = re.search(r'Packet Loss Rate:\s*([\d.]+)%', line)
            if loss_match:
                current_data['loss_rate'] = float(loss_match.group(1))

        elif line.startswith("RSSI:"):
            rssi_match = re.search(r'RSSI:\s*([\d.]+)\s*dBm', line)
            if rssi_match:
                current_data['rssi'] = float(rssi_match.group(1))

        elif line.startswith("Initial Phase:"):
            phase_match = re.search(r'Initial Phase:\s*([\d.]+)°', line)
            if phase_match:
                current_data['initial_phase'] = float(phase_match.group(1))

        elif line.startswith("Final Phase:"):
            phase_match = re.search(r'Final Phase:\s*([\d.]+)°', line)
            if phase_match:
                current_data['final_phase'] = float(phase_match.group(1))

        elif line.startswith("Reported Frequency:"):
            freq_match = re.search(r'Reported Frequency:\s*([\d.]+)\s*MHz', line)
            if freq_match:
                current_data['frequency'] = float(freq_match.group(1))

    # Don't forget to save the last tag
    if current_tag and current_data:
        freq = current_data.get('frequency')
        if freq:
            tag_data[current_tag][freq] = current_data.copy()

    return tag_data


def create_excel_output(tag_data, output_file):
    """
    Create Excel spreadsheet from parsed RFID data

    Args:
        tag_data (dict): Dictionary containing tag data organized by frequency
        output_file (str): Path for the output Excel file
    """
    # Get all unique frequencies
    all_frequencies = set()
    for tag_freqs in tag_data.values():
        all_frequencies.update(tag_freqs.keys())
    all_frequencies = sorted(all_frequencies)

    # Prepare data for Excel
    excel_data = []

    # Create header rows
    header1 = ['Tag EPC']
    header2 = ['']  # Empty cell below "Tag EPC"

    for freq in all_frequencies:
        header1.extend([f"{freq} MHz"] * 4)  # Each frequency spans 4 columns
        header2.extend(['RSSI (dBm)', 'Loss Rate (%)', 'Initial Phase (°)', 'Final Phase (°)'])

    excel_data.append(header1)
    excel_data.append(header2)

    # Add data rows for each tag
    for tag_epc, freq_data in tag_data.items():
        row = [tag_epc]

        for freq in all_frequencies:
            if freq in freq_data:
                data = freq_data[freq]
                row.extend([
                    data.get('rssi', ''),
                    data.get('loss_rate', ''),
                    data.get('initial_phase', ''),
                    data.get('final_phase', '')
                ])
            else:
                # No data for this frequency, add empty cells
                row.extend(['', '', '', ''])

        excel_data.append(row)

    # Create DataFrame and save to Excel
    df = pd.DataFrame(excel_data)

    # Save to Excel with formatting
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='RFID Data', index=False, header=False)

        # Get the workbook and worksheet for formatting
        workbook = writer.book
        worksheet = writer.sheets['RFID Data']

        # Apply formatting to header rows
        from openpyxl.styles import Font, PatternFill, Alignment

        # Format first header row (frequencies)
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color='E6E6FA', end_color='E6E6FA', fill_type='solid')

        # Format second header row (parameter names)
        param_font = Font(bold=True, italic=True)
        param_fill = PatternFill(start_color='F0F8FF', end_color='F0F8FF', fill_type='solid')

        # Apply formatting
        for row in range(1, 3):  # First two rows
            for col in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row, column=col)
                cell.alignment = Alignment(horizontal='center', vertical='center')

                if row == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                else:
                    cell.font = param_font
                    cell.fill = param_fill

        # Format tag names column
        for row in range(3, len(df) + 1):
            cell = worksheet.cell(row=row, column=1)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='left', vertical='center')

        # Auto-adjust column widths
        for column_cells in worksheet.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 20)


def main():
    """Main function to process RFID data"""
    if len(sys.argv) != 2:
        print("Usage: python process_rfid_data.py <input_file.txt>")
        print("Example: python process_rfid_data.py test.txt")
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = input_file.replace('.txt', '_processed.xlsx')

    try:
        print(f"Processing RFID data from: {input_file}")

        # Parse the data
        tag_data = parse_rfid_data(input_file)

        if not tag_data:
            print("No valid RFID data found in the file.")
            return

        print(f"Found data for {len(tag_data)} unique tags")

        # Count frequencies and show data distribution
        all_frequencies = set()
        for tag_epc, tag_freqs in tag_data.items():
            all_frequencies.update(tag_freqs.keys())
        print(f"Found {len(all_frequencies)} unique frequencies")

        # Show data distribution for verification
        print("\nData distribution per tag:")
        for tag_epc, tag_freqs in tag_data.items():
            print(f"  {tag_epc}: {len(tag_freqs)} frequency points")

        # Create Excel output
        create_excel_output(tag_data, output_file)

        print(f"Excel file created successfully: {output_file}")

    except FileNotFoundError:
        print(f"Error: File '{input_file}' not found.")
    except Exception as e:
        print(f"Error processing file: {e}")


if __name__ == "__main__":
    main()